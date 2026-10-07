"""Cálculo de indicadores e geração do relatório de vendas em Excel."""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def calculate_kpis(df: pd.DataFrame) -> dict:
    """Calcula os indicadores principais para o conjunto de vendas."""
    faturamento_total = float(df["total_value"].sum())
    quantidade_total = int(df["quantity"].sum())
    total_vendas = int(len(df))
    ticket_medio = faturamento_total / total_vendas if total_vendas else 0.0
    return {
        "faturamento_total": faturamento_total,
        "quantidade_total": quantidade_total,
        "total_vendas": total_vendas,
        "ticket_medio": ticket_medio,
    }


def _aggregate_sales(df: pd.DataFrame, group_column: str, include_quantity: bool = True) -> pd.DataFrame:
    aggregations = {"total_value": "sum"}
    if include_quantity:
        aggregations["quantity"] = "sum"
    grouped = df.groupby(group_column, as_index=False).agg(aggregations)
    grouped = grouped.rename(columns={"total_value": "faturamento_total"})
    if include_quantity:
        grouped = grouped.rename(columns={"quantity": "quantidade_total"})
    return grouped


def get_sales_by_category(df: pd.DataFrame) -> pd.DataFrame:
    """Agrupa faturamento, unidades e quantidade de vendas por categoria."""
    result = _aggregate_sales(df, "category")
    counts = df.groupby("category").size().rename("total_vendas")
    return result.join(counts, on="category")


def get_sales_by_seller(df: pd.DataFrame) -> pd.DataFrame:
    """Agrupa faturamento, unidades e quantidade de vendas por vendedor."""
    result = _aggregate_sales(df, "seller")
    counts = df.groupby("seller").size().rename("total_vendas")
    return result.join(counts, on="seller")


def get_sales_by_region(df: pd.DataFrame) -> pd.DataFrame:
    """Agrupa faturamento e unidades vendidas por região."""
    return _aggregate_sales(df, "region")


def get_sales_by_payment(df: pd.DataFrame) -> pd.DataFrame:
    """Agrupa faturamento e quantidade de vendas por forma de pagamento."""
    result = _aggregate_sales(df, "payment_method", include_quantity=False)
    counts = df.groupby("payment_method").size().rename("total_vendas")
    return result.join(counts, on="payment_method")


def get_monthly_trend(df: pd.DataFrame) -> pd.DataFrame:
    """Agrupa faturamento e unidades vendidas por mês (AAAA-MM)."""
    return _aggregate_sales(df, "year_month").sort_values("year_month")


def _format_workbook(writer: pd.ExcelWriter) -> None:
    """Aplica estilo de cabeçalho, filtros, congelamento e formatos numéricos."""
    currency_columns = {"faturamento_total", "ticket_medio"}
    integer_columns = {"quantidade_total", "total_vendas"}

    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="17365D")

        for column_cells in worksheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            header = str(column_cells[0].value)
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 14), 34)
            if header in currency_columns:
                for cell in column_cells[1:]:
                    cell.number_format = 'R$ #,##0.00'
            elif header in integer_columns:
                for cell in column_cells[1:]:
                    cell.number_format = "#,##0"


def generate_excel_report(df: pd.DataFrame, output_path: str | Path) -> None:
    """Escreve os KPIs e resumos em um arquivo Excel formatado."""
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    kpis = calculate_kpis(df)
    kpi_table = pd.DataFrame(
        [{"indicador": "Faturamento total", "valor": kpis["faturamento_total"]},
         {"indicador": "Quantidade total de itens", "valor": kpis["quantidade_total"]},
         {"indicador": "Total de vendas", "valor": kpis["total_vendas"]},
         {"indicador": "Ticket médio", "valor": kpis["ticket_medio"]}]
    )

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        kpi_table.to_excel(writer, sheet_name="Resumo KPIs", index=False)
        get_sales_by_category(df).to_excel(
            writer, sheet_name="Vendas por Categoria", index=False
        )
        get_sales_by_seller(df).to_excel(
            writer, sheet_name="Vendas por Vendedor", index=False
        )
        get_sales_by_region(df).to_excel(
            writer, sheet_name="Vendas por Região", index=False
        )
        get_sales_by_payment(df).to_excel(
            writer, sheet_name="Vendas por Pagamento", index=False
        )
        get_monthly_trend(df).to_excel(
            writer, sheet_name="Evolução Mensal", index=False
        )
        _format_workbook(writer)


def _format_brl(value: float) -> str:
    return f"R$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    data_path = project_root / "data" / "processed" / "vendas_processadas.csv"
    output_path = project_root / "output" / "relatorio_vendas.xlsx"

    if not data_path.is_file():
        raise FileNotFoundError(f"Dataset processado não encontrado: {data_path}")
    df = pd.read_csv(data_path)
    if df.empty:
        raise ValueError(f"Dataset processado está vazio: {data_path}")

    kpis = calculate_kpis(df)
    print("Resumo Executivo")
    print(f"Faturamento total: {_format_brl(kpis['faturamento_total'])}")
    print(f"Quantidade total de itens: {kpis['quantidade_total']:,}".replace(",", "."))
    print(f"Total de vendas: {kpis['total_vendas']:,}".replace(",", "."))
    print(f"Ticket médio: {_format_brl(kpis['ticket_medio'])}")

    generate_excel_report(df, output_path)
    print(f"Relatório Excel gerado: {output_path}")


if __name__ == "__main__":
    main()
